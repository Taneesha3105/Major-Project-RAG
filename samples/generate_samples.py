"""Generate a small fictional data room for demos and testing.

Run with the backend venv:  BE/.venv/bin/python samples/generate_samples.py
Companies and figures are fictional. Some inconsistencies are planted on purpose
(e.g. FY2024 revenue differs between documents) to exercise contradiction detection.
"""

from pathlib import Path

from docx import Document
from openpyxl import Workbook
from PIL import Image, ImageDraw, ImageFont

OUT = Path(__file__).parent / "data-room"
OUT.mkdir(exist_ok=True)

FINANCIALS = [
    "NIMBUS ANALYTICS LTD - AUDITED FINANCIAL STATEMENTS FY2024",
    "Independent auditor: Carrow & Pike LLP. Opinion: unqualified, with an emphasis of matter",
    "regarding the going concern assessment due to the revolving credit facility maturing in March 2026.",
    "",
    "INCOME STATEMENT (USD thousands)            FY2024      FY2023",
    "Revenue                                     48,200      41,750",
    "Cost of revenue                            (19,300)    (16,900)",
    "Gross profit                                28,900      24,850",
    "Operating expenses                         (22,400)    (19,100)",
    "EBITDA                                       9,850       8,120",
    "Net income                                   4,120       3,460",
    "",
    "BALANCE SHEET (USD thousands)",
    "Cash and equivalents                         6,300       7,900",
    "Trade receivables                           11,450       8,200",
    "Revolving credit facility (Harbor Bank)     18,000      12,000",
    "Total equity                                21,700      17,580",
    "",
    "Note 14 - Debt: The Harbor Bank facility contains a change of control clause permitting the",
    "lender to demand immediate repayment upon a change in ownership exceeding 50%.",
    "Covenant: net leverage must remain below 2.5x EBITDA, tested quarterly.",
    "",
    "Note 18 - Related parties: USD 1.2 million was paid to Larkspur Consulting, an entity",
    "owned by the CEO, for advisory services in FY2024.",
    "",
    "Note 21 - Contingencies: A former distributor, Velto GmbH, has filed a claim for",
    "EUR 3.4 million alleging wrongful termination. Management believes the claim is without",
    "merit and no provision has been recognised.",
]

MSA_PARAGRAPHS = [
    ("Heading 1", "Master Services Agreement"),
    (None, "This Master Services Agreement is entered into between Nimbus Analytics Ltd (\"Supplier\") "
           "and Brightwater Retail plc (\"Customer\") effective 1 January 2023 for an initial term of five years."),
    ("Heading 2", "1. Fees"),
    (None, "Customer shall pay annual subscription fees of USD 9,600,000, representing approximately 20% of "
           "Supplier's annual revenue. Fees increase by 3% annually."),
    ("Heading 2", "2. Termination"),
    (None, "Customer may terminate this Agreement for convenience on 90 days' written notice after 31 December 2025."),
    ("Heading 2", "3. Change of Control"),
    (None, "If Supplier undergoes a change of control, Customer may terminate this Agreement with immediate effect "
           "and shall be entitled to a refund of any prepaid fees. Supplier must notify Customer within 10 business "
           "days of any change of control."),
    ("Heading 2", "4. Limitation of Liability"),
    (None, "Supplier's aggregate liability is uncapped for data protection breaches and capped at 12 months' fees "
           "for all other claims."),
    ("Heading 2", "5. Exclusivity"),
    (None, "Supplier shall not provide competing analytics services to any UK grocery retailer during the term."),
]

MGMT_PRESENTATION = [
    "NIMBUS ANALYTICS - MANAGEMENT PRESENTATION (CONFIDENTIAL)",
    "Prepared for prospective acquirers - September 2025",
    "",
    "Highlights",
    "- FY2024 revenue of USD 52.0 million, up 25% year on year",
    "- 140 enterprise customers, top customer Brightwater Retail at 20% of revenue",
    "- Net revenue retention of 118%",
    "- Headcount: 212 employees across London, Pune and Austin",
    "",
    "Key people",
    "- The CEO and CTO hold 64% of the shares. Neither has a non-compete agreement.",
    "- Retention bonuses of USD 2.5 million are payable to 6 executives on a change of control.",
    "",
    "Tax",
    "- An HMRC enquiry into R&D tax credits claimed for FY2021-FY2022 (GBP 1.1 million) remains open.",
]


def _font(size: int):
    for candidate in ("/System/Library/Fonts/Supplemental/Arial.ttf", "/Library/Fonts/Arial.ttf",
                      "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"):
        if Path(candidate).exists():
            return ImageFont.truetype(candidate, size)
    return ImageFont.load_default()


def text_pdf(path: Path, lines: list[str]) -> None:
    """Minimal single-page PDF with a real text layer (no external dependencies)."""
    def esc(value: str) -> str:
        return value.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")

    stream_lines = ["BT", "/F1 9 Tf", "11 TL", "50 800 Td"]
    for line in lines:
        stream_lines.append(f"({esc(line)}) Tj T*")
    stream_lines.append("ET")
    stream = "\n".join(stream_lines).encode("latin-1")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Courier >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    path.write_bytes(bytes(out))


def scanned_pdf(path: Path, lines: list[str]) -> None:
    """Image-only PDF (no text layer) to simulate a scanned document that needs OCR."""
    image = Image.new("RGB", (1240, 1754), "white")
    draw = ImageDraw.Draw(image)
    font, title_font = _font(26), _font(34)
    y = 110
    for index, line in enumerate(lines):
        draw.text((90, y), line, fill=(20, 20, 20), font=title_font if index == 0 else font)
        y += 52 if index == 0 else 40
    image.save(path, "PDF", resolution=150)


def contract_docx(path: Path) -> None:
    document = Document()
    for style, text in MSA_PARAGRAPHS:
        if style:
            document.add_heading(text, level=1 if style == "Heading 1" else 2)
        else:
            document.add_paragraph(text)
    table = document.add_table(rows=3, cols=2)
    for row, (label, value) in zip(table.rows, [("Party", "Signatory"), ("Nimbus Analytics Ltd", "A. Mehta, CEO"),
                                                 ("Brightwater Retail plc", "J. Collins, COO")], strict=True):
        row.cells[0].text, row.cells[1].text = label, value
    document.save(path)


def customers_xlsx(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Top customers FY2024"
    sheet.append(["Customer", "Revenue (USD k)", "Share of revenue", "Contract end"])
    for row in [("Brightwater Retail plc", 9600, "19.9%", "2027-12-31"), ("Fenwick Logistics", 4100, "8.5%", "2026-06-30"),
                ("Osprey Insurance", 3350, "7.0%", "2025-12-31"), ("Calder Foods", 2900, "6.0%", "2028-03-31"),
                ("Marlow Health", 2200, "4.6%", "2026-09-30")]:
        sheet.append(list(row))
    workbook.save(path)


if __name__ == "__main__":
    text_pdf(OUT / "Nimbus_Audited_Financials_FY2024.pdf", FINANCIALS)
    scanned_pdf(OUT / "Nimbus_Management_Presentation_SCANNED.pdf", MGMT_PRESENTATION)
    contract_docx(OUT / "Brightwater_Master_Services_Agreement.docx")
    customers_xlsx(OUT / "Customer_Revenue_Analysis_FY2024.xlsx")
    print("Sample data room written to", OUT)
