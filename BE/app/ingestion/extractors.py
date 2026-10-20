"""Turn uploaded files into ordered text blocks with page/section provenance.

Native text is extracted with pdfplumber / python-docx / openpyxl. PDF pages (and images)
without a usable text layer are rendered and passed through OCR.
"""

from __future__ import annotations

import asyncio
import csv
import io
import logging
from dataclasses import dataclass, field
from pathlib import Path

from PIL import Image

from app.ai.ocr import ocr_image
from app.core.config import settings

logger = logging.getLogger(__name__)

PDF_EXTENSIONS = {".pdf"}
WORD_EXTENSIONS = {".docx"}
SHEET_EXTENSIONS = {".xlsx", ".xlsm"}
CSV_EXTENSIONS = {".csv"}
TEXT_EXTENSIONS = {".txt", ".md"}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp"}
SUPPORTED_EXTENSIONS = PDF_EXTENSIONS | WORD_EXTENSIONS | SHEET_EXTENSIONS | CSV_EXTENSIONS | TEXT_EXTENSIONS | IMAGE_EXTENSIONS


class ExtractionError(Exception):
    pass


@dataclass
class TextBlock:
    text: str
    page: int | None = None
    section: str | None = None
    is_ocr: bool = False
    is_heading: bool = False


@dataclass
class ExtractionResult:
    blocks: list[TextBlock] = field(default_factory=list)
    page_count: int | None = None
    ocr_pages: int = 0


async def extract(path: Path, extension: str) -> ExtractionResult:
    extension = extension.lower()
    if extension in PDF_EXTENSIONS:
        return await _extract_pdf(path)
    if extension in WORD_EXTENSIONS:
        return await asyncio.to_thread(_extract_docx, path)
    if extension in SHEET_EXTENSIONS:
        return await asyncio.to_thread(_extract_xlsx, path)
    if extension in CSV_EXTENSIONS:
        return await asyncio.to_thread(_extract_csv, path)
    if extension in TEXT_EXTENSIONS:
        return await asyncio.to_thread(_extract_text, path)
    if extension in IMAGE_EXTENSIONS:
        return await _extract_image(path)
    raise ExtractionError(f"Unsupported file type '{extension}'.")


# --- PDF -----------------------------------------------------------------------
def _pdf_text_layer(path: Path) -> list[str]:
    import pdfplumber

    pages: list[str] = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            try:
                pages.append(page.extract_text(x_tolerance=1.5, y_tolerance=3) or "")
            except Exception:  # noqa: BLE001 - a single corrupt page should not fail the file
                logger.warning("pdfplumber failed on page %s of %s", page.page_number, path.name)
                pages.append("")
            finally:
                page.close()
    return pages


def _render_pdf_page(path: Path, index: int) -> Image.Image:
    import pypdfium2 as pdfium

    pdf = pdfium.PdfDocument(path)
    try:
        page = pdf[index]
        image = page.render(scale=settings.ocr_render_scale).to_pil()
        page.close()
        return image
    finally:
        pdf.close()


async def _extract_pdf(path: Path) -> ExtractionResult:
    try:
        page_texts = await asyncio.to_thread(_pdf_text_layer, path)
    except Exception as exc:  # noqa: BLE001
        raise ExtractionError("The PDF could not be opened. It may be encrypted or corrupted.") from exc

    result = ExtractionResult(page_count=len(page_texts))
    for index, text in enumerate(page_texts):
        page_number = index + 1
        if settings.ocr_enabled and len(text.strip()) < settings.ocr_min_chars_per_page:
            image = await asyncio.to_thread(_render_pdf_page, path, index)
            ocr_text = await ocr_image(image)
            if ocr_text.strip():
                result.ocr_pages += 1
                result.blocks.append(TextBlock(ocr_text, page=page_number, is_ocr=True))
                continue
        if text.strip():
            result.blocks.append(TextBlock(text, page=page_number))
    return result


# --- Images --------------------------------------------------------------------
async def _extract_image(path: Path) -> ExtractionResult:
    if not settings.ocr_enabled:
        raise ExtractionError("OCR is disabled, so image files cannot be processed.")
    try:
        image = Image.open(path)
    except Exception as exc:  # noqa: BLE001
        raise ExtractionError("The image could not be opened.") from exc

    frames = getattr(image, "n_frames", 1)
    result = ExtractionResult(page_count=frames)
    for frame in range(frames):
        image.seek(frame)
        text = await ocr_image(image.copy())
        if text.strip():
            result.ocr_pages += 1
            result.blocks.append(TextBlock(text, page=frame + 1, is_ocr=True))
    return result


# --- Word ----------------------------------------------------------------------
def _extract_docx(path: Path) -> ExtractionResult:
    from docx import Document as DocxDocument
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    try:
        document = DocxDocument(str(path))
    except Exception as exc:  # noqa: BLE001
        raise ExtractionError("The Word document could not be opened.") from exc

    result = ExtractionResult()
    section: str | None = None
    for element in document.element.body.iterchildren():
        tag = element.tag.rsplit("}", 1)[-1]
        if tag == "p":
            paragraph = Paragraph(element, document)
            text = paragraph.text.strip()
            if not text:
                continue
            style = (paragraph.style.name or "").lower() if paragraph.style is not None else ""
            if style.startswith("heading") or style == "title":
                section = text[:300]
                result.blocks.append(TextBlock(text, section=section, is_heading=True))
            else:
                result.blocks.append(TextBlock(text, section=section))
        elif tag == "tbl":
            rows = []
            for row in Table(element, document).rows:
                cells: list[str] = []
                for cell in row.cells:
                    value = " ".join(cell.text.split())
                    if not cells or cells[-1] != value:  # merged cells repeat their text
                        cells.append(value)
                if any(cells):
                    rows.append(" | ".join(cells))
            if rows:
                result.blocks.append(TextBlock("\n".join(rows), section=section))
    return result


# --- Spreadsheets --------------------------------------------------------------
def _format_cell(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float):
        return f"{value:,.2f}".rstrip("0").rstrip(".")
    return str(value).strip()


def _rows_to_blocks(rows: list[list[str]], *, page: int, section: str, rows_per_block: int = 40) -> list[TextBlock]:
    """Repeat the header row in each block so every chunk keeps its column meaning."""
    rows = [row for row in rows if any(cell for cell in row)]
    if not rows:
        return []
    header, body = rows[0], rows[1:] or []
    blocks = []
    if not body:
        return [TextBlock(" | ".join(header), page=page, section=section)]
    for start in range(0, len(body), rows_per_block):
        lines = [" | ".join(header)] + [" | ".join(row) for row in body[start : start + rows_per_block]]
        blocks.append(TextBlock("\n".join(lines), page=page, section=section))
    return blocks


def _extract_xlsx(path: Path) -> ExtractionResult:
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001
        raise ExtractionError("The spreadsheet could not be opened.") from exc

    result = ExtractionResult(page_count=len(workbook.sheetnames))
    try:
        for index, sheet in enumerate(workbook.worksheets, start=1):
            rows = [[_format_cell(cell) for cell in row] for row in sheet.iter_rows(values_only=True)]
            result.blocks.extend(_rows_to_blocks(rows, page=index, section=f"Sheet: {sheet.title}"))
    finally:
        workbook.close()
    return result


def _read_text_file(path: Path) -> str:
    raw = path.read_bytes()
    for encoding in ("utf-8-sig", "utf-16", "cp1252", "latin-1"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="ignore")


def _extract_csv(path: Path) -> ExtractionResult:
    text = _read_text_file(path)
    try:
        dialect = csv.Sniffer().sniff(text[:4096])
    except csv.Error:
        dialect = csv.excel
    rows = [[cell.strip() for cell in row] for row in csv.reader(io.StringIO(text), dialect)]
    return ExtractionResult(blocks=_rows_to_blocks(rows, page=1, section=f"Table: {path.stem}"), page_count=1)


def _extract_text(path: Path) -> ExtractionResult:
    result = ExtractionResult()
    section: str | None = None
    for paragraph in _read_text_file(path).split("\n\n"):
        text = paragraph.strip()
        if not text:
            continue
        if text.startswith("#"):
            heading = text.splitlines()[0].lstrip("#").strip()
            section = heading[:300] or section
        result.blocks.append(TextBlock(text, section=section))
    return result
