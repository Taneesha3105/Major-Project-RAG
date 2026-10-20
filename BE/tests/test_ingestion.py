import asyncio
from pathlib import Path

from docx import Document as DocxDocument
from openpyxl import Workbook

from app.ingestion.chunker import chunk_blocks
from app.ingestion.cleaning import clean_blocks, normalize_text
from app.ingestion.extractors import TextBlock, extract


def test_normalize_text_fixes_hyphenation_and_whitespace():
    assert normalize_text("Reve-\nnue   grew\t\t12%\n\n\n\nNext") == "Revenue grew 12%\n\nNext"


def test_repeated_headers_and_footers_are_removed():
    blocks = [TextBlock(f"CONFIDENTIAL\nBody text of page {p} with facts.\nPage {p} of 5", page=p) for p in range(1, 6)]
    cleaned = clean_blocks(blocks)
    assert all("CONFIDENTIAL" not in b.text and "of 5" not in b.text for b in cleaned)
    assert all("Body text" in b.text for b in cleaned)


def test_chunks_respect_size_pages_and_overlap():
    text = " ".join(f"Clause {n} of the facility allows early repayment." for n in range(60))
    blocks = [TextBlock(text, page=1), TextBlock("Short page two text that is long enough.", page=2)]
    chunks = chunk_blocks(blocks, chunk_size=500, overlap=120)
    assert all(len(c.content) <= 500 for c in chunks)
    assert {c.page for c in chunks} == {1, 2}
    page_one = [c for c in chunks if c.page == 1]
    assert len(page_one) > 2
    # the next chunk starts with the last sentence(s) of the previous one
    last_sentence = page_one[0].content.rsplit("Clause ", 1)[1]
    assert page_one[1].content.startswith("Clause " + last_sentence) or ("Clause " + last_sentence) in page_one[1].content[:150]


def test_headings_set_section_metadata():
    blocks = [
        TextBlock("1. Fees", section="1. Fees", is_heading=True),
        TextBlock("Annual fees are USD 9.6m." * 3, section="1. Fees"),
        TextBlock("2. Termination", section="2. Termination", is_heading=True),
        TextBlock("Either party may terminate on 90 days notice." * 3, section="2. Termination"),
    ]
    chunks = chunk_blocks(blocks, chunk_size=200, overlap=0)
    assert any(c.section == "2. Termination" and "terminate" in c.content for c in chunks)


def test_extract_docx_and_xlsx(tmp_path: Path):
    docx_path = tmp_path / "contract.docx"
    document = DocxDocument()
    document.add_heading("Change of Control", level=1)
    document.add_paragraph("Customer may terminate upon a change of control.")
    table = document.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Party", "Role"
    table.rows[1].cells[0].text, table.rows[1].cells[1].text = "Nimbus", "Supplier"
    document.save(docx_path)

    result = asyncio.run(extract(docx_path, ".docx"))
    texts = [b.text for b in result.blocks]
    assert "Change of Control" in texts[0] and result.blocks[0].is_heading
    assert any("Nimbus | Supplier" in t for t in texts)
    assert result.blocks[1].section == "Change of Control"

    xlsx_path = tmp_path / "customers.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Top"
    sheet.append(["Customer", "Revenue"])
    sheet.append(["Brightwater", 9600.5])
    workbook.save(xlsx_path)

    result = asyncio.run(extract(xlsx_path, ".xlsx"))
    assert result.blocks[0].section == "Sheet: Top"
    assert "Customer | Revenue\nBrightwater | 9,600.5" == result.blocks[0].text
